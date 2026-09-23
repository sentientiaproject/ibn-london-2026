import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_strategy_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    NAVY_DARK = RGBColor(6, 18, 38)       # #061226
    NAVY_CARD = RGBColor(12, 28, 54)      # #0C1C36
    NAVY_LIGHT = RGBColor(18, 38, 72)     # #122648
    GOLD = RGBColor(212, 175, 55)         # #D4AF37
    GOLD_LIGHT = RGBColor(243, 229, 171)  # #F3E5AB
    WHITE = RGBColor(255, 255, 255)
    MUTED_GRAY = RGBColor(180, 195, 215)
    BORDER_COLOR = RGBColor(30, 58, 95)

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="GLOBAL INDIAN BUSINESS EXCELLENCE AWARDS 2026"):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Georgia"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = GOLD
        p_cat.space_after = Pt(4)

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.name = "Calibri"
        p_title.font.size = Pt(25)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.48), Inches(1.8), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = GOLD
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=NAVY_CARD, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.color.rgb = BORDER_COLOR
            card.line.width = Pt(1)
        return card

    # ==========================================
    # SLIDE 1: Title & Cover
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    outer = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.4), Inches(12.533), Inches(6.7))
    outer.fill.background()
    outer.line.color.rgb = RGBColor(40, 70, 110)
    outer.line.width = Pt(1)

    inner = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.45), Inches(12.433), Inches(6.6))
    inner.fill.background()
    inner.line.color.rgb = GOLD
    inner.line.width = Pt(0.75)

    if os.path.exists("presentation_assets/big_ben_tower.png"):
        s1.shapes.add_picture("presentation_assets/big_ben_tower.png", Inches(8.4), Inches(0.55), Inches(4.3), Inches(6.4))

    if os.path.exists("presentation_assets/ibn_logo_top.png"):
        s1.shapes.add_picture("presentation_assets/ibn_logo_top.png", Inches(1.0), Inches(0.8), Inches(3.8), Inches(0.75))

    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(7.2), Inches(3.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "INDIAN BUSINESS NETWORK PRESENTS"
    p0.font.name = "Georgia"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = GOLD
    p0.space_after = Pt(10)

    p1 = tf.add_paragraph()
    p1.text = "OMNICHANNEL MARKETING &\nDIGITAL PLATFORM STRATEGY"
    p1.font.name = "Georgia"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = GOLD_LIGHT
    p1.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "Mobile App • Website • LinkedIn • Facebook • Instagram • WhatsApp • Email • Print Collateral\nHouse of Commons, British Parliament • Oxford • Cambridge • Imperial College London"
    p2.font.name = "Calibri"
    p2.font.size = Pt(13)
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(20)

    pill = add_card(s1, Inches(1.0), Inches(5.1), Inches(7.0), Inches(1.2), NAVY_CARD, GOLD)
    ptb = s1.shapes.add_textbox(Inches(1.2), Inches(5.2), Inches(6.6), Inches(1.0))
    ptf = ptb.text_frame
    pp = ptf.paragraphs[0]
    pp.text = "🏛️ 5–9 NOVEMBER 2026  •  🔒 STRICTLY LIMITED TO 100 GUESTS"
    pp.font.name = "Calibri"
    pp.font.size = Pt(13)
    pp.font.bold = True
    pp.font.color.rgb = GOLD

    pp2 = ptf.add_paragraph()
    pp2.text = "Digital Platform Strategy • Print Collateral • Onsite Event Distribution"
    pp2.font.name = "Calibri"
    pp2.font.size = Pt(11.5)
    pp2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 2: Mobile App Strategy (iOS & Android)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "Mobile App Strategy: iOS & Android VIP Event Companion", "MOBILE APP STRATEGY (SCOPE ONLY)")

    app_cards = [
        ("presentation_assets/badges/shield_security.png", "Parliament Security KYC", "Pre-event identity and passport verification gateway to ensure seamless security clearance at the House of Commons Cromwell Green entrance.", GOLD),
        ("presentation_assets/badges/mobile_wallet.png", "Dynamic Pass & Mobile Wallet", "Anti-screenshot dynamic QR badge integrated with Apple Wallet and Google Wallet, automatically displaying near Westminster.", WHITE),
        ("presentation_assets/badges/coach_transit.png", "5-Day Logistics & Transit GPS", "Hour-by-hour interactive agenda, dress code advisories, and live luxury coach tracking for transfers to Oxford and Cambridge.", GOLD_LIGHT),
        ("presentation_assets/badges/payment_card.png", "Compliant B2B Billing", "Multi-currency checkout (GBP, USD, EUR, INR) with instant corporate tax invoices for leadership programme fees.", WHITE),
        ("presentation_assets/badges/match_meeting.png", "1-on-1 Bilateral Matchmaker", "Curated directory of 100 leaders with built-in meeting scheduler to book 15-min private catch-ups during breaks.", GOLD),
        ("presentation_assets/icons/icon_photo.png", "AI Facial-Recognition Photos", "Delegates upload a selfie; the app automatically curates all official high-res photos they appear in for immediate download.", WHITE),
    ]

    for m, (ic_path, a_title, a_desc, a_col) in enumerate(app_cards):
        row = m // 3
        col = m % 3
        bx = Inches(0.8 + col * 4.04)
        by = Inches(1.8 + row * 2.5)
        add_card(s2, bx, by, Inches(3.64), Inches(2.2), NAVY_CARD, GOLD if m==0 else None)

        if os.path.exists(ic_path):
            s2.shapes.add_picture(ic_path, bx + Inches(0.2), by + Inches(0.2), Inches(0.55), Inches(0.55))

        box = s2.shapes.add_textbox(bx + Inches(0.85), by + Inches(0.15), Inches(2.65), Inches(1.9))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = a_title
        p1.font.name = "Georgia"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = a_col
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = a_desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 3: Website Strategy (Scope & Purpose)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Official Website Strategy: Digital Flagship & Vetting Hub", "WEBSITE STRATEGY (SCOPE ONLY)")

    web_cards = [
        ("presentation_assets/badges/executive_vip.png", "Sovereign Digital Flagship", "Validates institutional prestige, celebrating Indian business impact at the House of Commons, Oxford, Cambridge, and Imperial College London.", GOLD),
        ("presentation_assets/badges/visa_vault.png", "Delegate Vetting & Application", "Multi-step Expression of Interest (EOI) portal (Tally integrated) to filter and select exactly 100 qualified C-Suite delegates.", WHITE),
        ("presentation_assets/badges/coach_transit.png", "5-Day Interactive Timeline", "Tabbed journey guide providing deep transparency on session themes, academic forums, dress codes, and logistics.", GOLD_LIGHT),
        ("presentation_assets/badges/shield_security.png", "Statutory & Legal Separation", "Clear regulatory notice establishing that awards attendance is by invitation only and fees solely fund the business leadership programme.", WHITE),
        ("presentation_assets/icons/icon_cert.png", "Award Categories & Nominations", "Showcases 8 core recognition categories with criteria, enabling peer-to-peer nomination and credentials submission.", GOLD),
        ("presentation_assets/badges/scanner_gate.png", "Concierge & FAQ Helpdesk", "Interactive answers for UK visa support letters, hotel reservations, airport chauffeur bookings, and secretariat hotlines.", WHITE),
    ]

    for q, (ic_path, w_title, w_desc, w_col) in enumerate(web_cards):
        row = q // 3
        col = q % 3
        wx = Inches(0.8 + col * 4.04)
        wy = Inches(1.8 + row * 2.5)
        add_card(s3, wx, wy, Inches(3.64), Inches(2.2), NAVY_CARD, GOLD if q==1 else None)

        if os.path.exists(ic_path):
            s3.shapes.add_picture(ic_path, wx + Inches(0.2), wy + Inches(0.2), Inches(0.55), Inches(0.55))

        box = s3.shapes.add_textbox(wx + Inches(0.85), wy + Inches(0.15), Inches(2.65), Inches(1.9))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = w_title
        p1.font.name = "Georgia"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = w_col
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = w_desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 4: Printable Brochure & Onsite Template Distribution
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Printable Brochure & Onsite Event Template Distribution", "EVENT COLLATERAL & DISTRIBUTION")

    print_cards = [
        ("presentation_assets/icons/icon_kit.png", "Print-Ready Executive Brochure", "Luxury multi-page commemorative brochure formatted in high-resolution CMYK with gold-foil finish specs. Available as a downloadable print PDF and placed inside the physical delegate welcome briefcase.", GOLD),
        ("presentation_assets/icons/icon_cert.png", "Award Winner Press Release Kits", "Pre-formatted press release templates and high-res digital award crests distributed to winners to instantly syndicate to international news agencies (Bloomberg, PTI, Financial Times, Reuters).", WHITE),
        ("presentation_assets/badges/match_meeting.png", "Bilateral MoU & Deal Templates", "Standardized bilateral business cooperation and investment intent (MoU) templates provided during the Oxford & Imperial forums to facilitate immediate partnership signings.", GOLD_LIGHT),
        ("presentation_assets/icons/icon_badge.png", "Onsite Delegate Stationery & Badges", "Custom embossed executive lanyards, personalized summit itinerary handbooks, table seating cards, and formal embossed Participation Certificates presented on Day 5.", WHITE),
    ]

    for p_idx, (ic_path, p_head, p_body, p_col) in enumerate(print_cards):
        row = p_idx // 2
        col = p_idx % 2
        px = Inches(0.8 + col * 6.04)
        py = Inches(1.8 + row * 2.5)
        add_card(s4, px, py, Inches(5.68), Inches(2.2), NAVY_CARD, GOLD if p_idx==0 else None)

        if os.path.exists(ic_path):
            s4.shapes.add_picture(ic_path, px + Inches(0.25), py + Inches(0.25), Inches(0.65), Inches(0.65))

        box = s4.shapes.add_textbox(px + Inches(1.05), py + Inches(0.2), Inches(4.4), Inches(1.8))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "✦ " + p_head
        p1.font.name = "Georgia"
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = p_col
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = p_body
        p2.font.name = "Calibri"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 5: LinkedIn Marketing Strategy
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "LinkedIn Promotion: C-Suite Authority & Thought Leadership", "MARKETING STRATEGY: LINKEDIN")

    li_left = add_card(s5, Inches(0.8), Inches(1.8), Inches(7.2), Inches(4.9), NAVY_CARD, GOLD)
    lbox = s5.shapes.add_textbox(Inches(1.1), Inches(2.05), Inches(6.6), Inches(4.4))
    ltf = lbox.text_frame
    ltf.word_wrap = True

    lp1 = ltf.paragraphs[0]
    lp1.text = "CORE MARKETING ENGINE: B2B & EXECUTIVE VISIBILITY"
    lp1.font.name = "Georgia"
    lp1.font.size = Pt(14)
    lp1.font.bold = True
    lp1.font.color.rgb = GOLD
    lp1.space_after = Pt(10)

    li_pillars = [
        ("Institutional Announcement", "Positioning the summit as an exclusive UK–India bilateral leadership summit held inside the British Parliament."),
        ("Founder & President Voice", "Personal executive notes from Velou Singaram and board members inviting peers to join the private corridor."),
        ("Academic Immersion Highlights", "Spotlighting Oxford, Cambridge, and Imperial College London as catalysts for deep-tech and enterprise scaling."),
        ("Delegate & Speaker Spotlights", "Elegant digital badge creatives announcing vetted attendees: 'Proud to announce [Name] as an Executive Delegate.'"),
        ("100-Seat Urgency Triggers", "Security clearance deadline countdowns driving C-Suite leaders to submit applications before the quota closes.")
    ]
    for p_head, p_body in li_pillars:
        p = ltf.add_paragraph()
        p.text = f"• {p_head}: {p_body}"
        p.font.name = "Calibri"
        p.font.size = Pt(11.5)
        p.font.color.rgb = WHITE
        p.space_after = Pt(6)

    li_right_items = [
        ("TARGET AUDIENCE", "Indian Founders, CEOs, Investors (UK, India, UAE, US, Europe)"),
        ("HASHTAG STRATEGY", "#GlobalIndianAwards #HouseOfCommons #UKIndiaTrade #IBN2026"),
        ("CONVERSION HOOK", "Direct link to Tally Vetting Form with VIP nomination incentive"),
    ]
    for r_idx, (r_title, r_desc) in enumerate(li_right_items):
        ry = Inches(1.8 + r_idx * 1.68)
        add_card(s5, Inches(8.3), ry, Inches(4.233), Inches(1.5), NAVY_LIGHT)
        box = s5.shapes.add_textbox(Inches(8.5), ry + Inches(0.2), Inches(3.8), Inches(1.1))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = r_title
        p1.font.name = "Georgia"
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = GOLD

        p2 = tf.add_paragraph()
        p2.text = r_desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 6: Facebook Marketing Strategy
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Facebook Promotion: Diaspora Communities & Multi-Card Carousels", "MARKETING STRATEGY: FACEBOOK")

    fb_cards = [
        ("Official Event Page Hub", "Create a verified Facebook Event Page: 'Global Indian Business Excellence Awards 2026 — London', serving as the primary RSVP and live announcement center.", GOLD),
        ("Multi-Slide Carousel Ads", "'5 Days. 3 Historic Institutions. 100 Leaders.' Multi-card ads showcasing Big Ben, the House of Commons, Gala ballroom, Oxford, and Cambridge.", WHITE),
        ("Business Diaspora Groups", "Engage established British-Indian diaspora groups, executive alumni networks (IIT/IIM UK Chapters), and bilateral business councils.", GOLD_LIGHT),
        ("Video Teaser Campaigns", "Short 30-second cinematic teasers highlighting the prestige of Westminster and academic innovation hubs, targeted to mature business demographics.", WHITE),
    ]

    for f_idx, (f_head, f_body, f_col) in enumerate(fb_cards):
        row = f_idx // 2
        col = f_idx % 2
        fx = Inches(0.8 + col * 6.04)
        fy = Inches(1.8 + row * 2.5)
        add_card(s6, fx, fy, Inches(5.68), Inches(2.2), NAVY_CARD, GOLD if f_idx==1 else None)

        box = s6.shapes.add_textbox(fx + Inches(0.3), fy + Inches(0.25), Inches(5.08), Inches(1.7))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "✦ " + f_head
        p1.font.name = "Georgia"
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = f_col
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = f_body
        p2.font.name = "Calibri"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 7: Instagram Marketing Strategy
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "Instagram Promotion: Visual Prestige, Cinematic Reels & Stories", "MARKETING STRATEGY: INSTAGRAM")

    ig_cards = [
        ("Cinematic Luxury Reels (9:16)", "35-second high-energy cinematic clips: Big Ben twilight aerials, red-carpet Black Tie gala, Oxford colleges, and the Parliament terrace.", GOLD),
        ("High-Engagement Feed Carousels", "10-slide luxury slide carousels breaking down each day's itinerary, delegate package inclusions, and award categories with dark navy/gold aesthetics.", WHITE),
        ("Interactive Countdown Stories", "Countdown stickers for registration closing, 'Meet the Committee' story cards, dress-code style inspiration (Black Tie / Tuxedo), and live Q&A stickers.", GOLD_LIGHT),
        ("Curated Highlight Covers", "Permanent profile highlights on Instagram: 🏛️ Parliament | 🗓️ Itinerary | 🎓 Oxford & Cambridge | 👔 Dress Code | 🎟️ Apply.", WHITE),
    ]

    for i_idx, (i_head, i_body, i_col) in enumerate(ig_cards):
        row = i_idx // 2
        col = i_idx % 2
        ix = Inches(0.8 + col * 6.04)
        iy = Inches(1.8 + row * 2.5)
        add_card(s7, ix, iy, Inches(5.68), Inches(2.2), NAVY_CARD, GOLD if i_idx==0 else None)

        box = s7.shapes.add_textbox(ix + Inches(0.3), iy + Inches(0.25), Inches(5.08), Inches(1.7))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "✦ " + i_head
        p1.font.name = "Georgia"
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = i_col
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = i_body
        p2.font.name = "Calibri"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 8: WhatsApp Marketing & Concierge
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "WhatsApp Strategy: High-Touch VIP Concierge & Direct Outreach", "MARKETING STRATEGY: WHATSAPP")

    wa_cards = [
        ("Personalized 1-on-1 VIP Outreach", "Direct confidential invitations sent to nominated CEOs and industry titans with personalized PDF invitations and brochure attachments.", GOLD),
        ("Instant Secretariat Hotline", "One-tap direct chat with Event Coordinator Subha Austalekshmi (+44 7587 260254) for immediate answers regarding itinerary, hotels, and visas.", WHITE),
        ("Accredited Delegate Broadcast Channel", "Private, announcement-only broadcast group for the 100 confirmed delegates providing daily security timing and coach departure updates.", GOLD_LIGHT),
        ("Interactive WhatsApp Chatbot Flow", "Automated FAQ assistant answering queries on dress code, Cromwell Green passport requirements, and coach pickup points in central London.", WHITE),
    ]

    for w_idx, (w_head, w_body, w_col) in enumerate(wa_cards):
        row = w_idx // 2
        col = w_idx % 2
        wx = Inches(0.8 + col * 6.04)
        wy = Inches(1.8 + row * 2.5)
        add_card(s8, wx, wy, Inches(5.68), Inches(2.2), NAVY_CARD, GOLD if w_idx==1 else None)

        box = s8.shapes.add_textbox(wx + Inches(0.3), wy + Inches(0.25), Inches(5.08), Inches(1.7))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "✦ " + w_head
        p1.font.name = "Georgia"
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = w_col
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = w_body
        p2.font.name = "Calibri"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 9: Email Newsletter Drip Sequence
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9)
    add_header(s9, "Email Newsletter Campaign: 4-Stage High-Conversion VIP Drip", "MARKETING STRATEGY: EMAIL")

    emails = [
        ("EMAIL 1: VIP INVITATION", "The British Parliament Calling", "Formal high-touch invitation addressed directly to C-Suite leaders detailing the 5-day summit, House of Commons prestige, and the 100-seat ceiling.", GOLD),
        ("EMAIL 2: ACADEMIC FOCUS", "Beyond Westminster: Oxford & Cambridge", "Deep dive into the Oxford Business Leadership Forum, Cambridge deep-tech cluster, and Imperial College London research commercialisation.", WHITE),
        ("EMAIL 3: URGENCY & CLEARANCE", "Capacity Warning & Security Lead Times", "Alerting invited leaders that 80% of capacity is allocated and parliamentary accreditation requires passport details well in advance.", GOLD_LIGHT),
        ("EMAIL 4: ONBOARDING BRIEFING", "Confirmed Delegate Welcome Package", "Dispatched to confirmed leaders with dress codes (Black Tie / Formal), security clearance confirmations, and mobile app download credentials.", WHITE),
    ]

    for e_idx, (e_phase, e_title, e_desc, e_col) in enumerate(emails):
        ex = Inches(0.8 + e_idx * 3.03)
        ey = Inches(1.8)
        add_card(s9, ex, ey, Inches(2.83), Inches(4.9), NAVY_CARD, GOLD if e_idx==0 else None)

        box = s9.shapes.add_textbox(ex + Inches(0.2), ey + Inches(0.3), Inches(2.43), Inches(4.3))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = e_phase
        p1.font.name = "Georgia"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = e_col
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = e_title
        p2.font.name = "Calibri"
        p2.font.size = Pt(15)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_after = Pt(12)

        p3 = tf.add_paragraph()
        p3.text = e_desc
        p3.font.name = "Calibri"
        p3.font.size = Pt(11.5)
        p3.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 10: Integrated Marketing Funnel
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10)
    add_header(s10, "Integrated Multi-Channel Conversion Funnel", "MARKETING EXECUTION")

    funnel_stages = [
        ("STAGE 1: AWARENESS", "LinkedIn & Media PR", "Thought leadership articles, official press releases, and executive announcements establishing global prestige."),
        ("STAGE 2: ENGAGEMENT", "Instagram & Facebook", "Cinematic video reels, carousel ads, and community discussions highlighting Oxford, Cambridge, and the Gala."),
        ("STAGE 3: CONVERSION", "Website & WhatsApp", "Official vetting portal on the website (Tally embed) supported by 1-on-1 WhatsApp VIP concierge consultations."),
        ("STAGE 4: NURTURING", "Email Drip & Mobile App", "Targeted 4-stage email newsletters, security vetting confirmations, and digital mobile pass activation."),
    ]

    for s_idx, (f_stage, f_chan, f_desc) in enumerate(funnel_stages):
        sx = Inches(0.8 + s_idx * 3.03)
        sy = Inches(1.8)
        add_card(s10, sx, sy, Inches(2.83), Inches(4.9), NAVY_CARD, GOLD if s_idx==2 else None)

        box = s10.shapes.add_textbox(sx + Inches(0.2), sy + Inches(0.3), Inches(2.43), Inches(4.3))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = f_stage
        p1.font.name = "Georgia"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = GOLD
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = f_chan
        p2.font.name = "Calibri"
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_after = Pt(12)

        p3 = tf.add_paragraph()
        p3.text = f_desc
        p3.font.name = "Calibri"
        p3.font.size = Pt(12)
        p3.font.color.rgb = MUTED_GRAY

    output_path = "IBN_2026_Digital_Marketing_and_Platform_Strategy.pptx"
    prs.save(output_path)
    print(f"Strategy presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_strategy_deck()
