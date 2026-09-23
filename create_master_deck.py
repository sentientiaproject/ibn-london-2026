import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_master_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    NAVY_DARK = RGBColor(6, 18, 38)       # #061226
    NAVY_CARD = RGBColor(12, 28, 54)      # #0C1C36
    NAVY_LIGHT = RGBColor(20, 42, 78)     # #142A4E
    GOLD = RGBColor(212, 175, 55)         # #D4AF37
    GOLD_LIGHT = RGBColor(243, 229, 171)  # #F3E5AB
    WHITE = RGBColor(255, 255, 255)
    MUTED_GRAY = RGBColor(180, 195, 215)
    BORDER_COLOR = RGBColor(30, 58, 95)

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="GLOBAL INDIAN BUSINESS EXCELLENCE AWARDS 2026"):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Georgia"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = GOLD
        p_cat.space_after = Pt(4)

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.name = "Calibri"
        p_title.font.size = Pt(25)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.48), Inches(1.8), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = GOLD
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=NAVY_CARD, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.color.rgb = BORDER_COLOR
            card.line.width = Pt(1)
        return card

    # ==========================================
    # SLIDE 1: Title & Cover
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    outer = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.4), Inches(12.533), Inches(6.7))
    outer.fill.background()
    outer.line.color.rgb = RGBColor(40, 70, 110)
    outer.line.width = Pt(1)

    inner = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.45), Inches(12.433), Inches(6.6))
    inner.fill.background()
    inner.line.color.rgb = GOLD
    inner.line.width = Pt(0.75)

    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    p0.text = "INDIAN BUSINESS NETWORK PRESENTS"
    p0.font.name = "Georgia"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = GOLD
    p0.space_after = Pt(15)

    p1 = tf.add_paragraph()
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "GLOBAL INDIAN BUSINESS EXCELLENCE AWARDS 2026\nEVENT MOBILE APPLICATION: FEATURES & ARCHITECTURE"
    p1.font.name = "Georgia"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = GOLD_LIGHT
    p1.space_after = Pt(15)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "A Purpose-Built Executive Platform for 100 Global Business Leaders\nHouse of Commons, British Parliament • Oxford • Cambridge • Imperial College London"
    p2.font.name = "Calibri"
    p2.font.size = Pt(15)
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(25)

    pill = add_card(s1, Inches(2.6), Inches(4.7), Inches(8.133), Inches(1.1), NAVY_CARD, GOLD)
    ptb = s1.shapes.add_textbox(Inches(2.8), Inches(4.85), Inches(7.733), Inches(0.8))
    ptf = ptb.text_frame
    pp = ptf.paragraphs[0]
    pp.alignment = PP_ALIGN.CENTER
    pp.text = "🏛️ 5–9 NOVEMBER 2026  •  📱 iOS & ANDROID NATIVE COMPANION"
    pp.font.name = "Calibri"
    pp.font.size = Pt(13)
    pp.font.bold = True
    pp.font.color.rgb = GOLD

    pp2 = ptf.add_paragraph()
    pp2.alignment = PP_ALIGN.CENTER
    pp2.text = "Parliament Security Clearance • 5-Day Logistics • Multi-Currency Payments • AI Matchmaking"
    pp2.font.name = "Calibri"
    pp2.font.size = Pt(12)
    pp2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 2: Executive Vision & Context
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "Executive Rationale: Purpose-Built for UHNW Leadership")

    pillars = [
        ("01. DIPLOMATIC PROTOCOL", "British Parliament Access", "The House of Commons requires rigorous advance identity accreditation. The platform secures passport KYC data and delivers anti-counterfeit digital passes.", GOLD),
        ("02. MULTI-CITY TRANSIT", "London, Oxford & Cambridge", "Coordinates 100 VIP delegates across 3 world-renowned academic cities over 5 days, featuring coach GPS tracking and real-time itinerary prompts.", WHITE),
        ("03. C-SUITE NETWORKING", "Privacy-Preserved Corridors", "Empowers high-level leaders to connect, chat, and schedule bilateral 1-on-1 meetings without exposing private phone numbers or personal emails.", GOLD_LIGHT),
        ("04. COMPLIANT BILLING", "Transparent Package Invoicing", "Automates multi-currency payments and corporate B2B invoicing while maintaining the non-monetary invitation status of the Parliamentary ceremony.", WHITE),
    ]

    for i, (tag, head, desc, col) in enumerate(pillars):
        x = Inches(0.8 + i * 3.03)
        y = Inches(1.8)
        add_card(s2, x, y, Inches(2.83), Inches(4.9), NAVY_CARD, GOLD if i==0 else None)

        box = s2.shapes.add_textbox(x + Inches(0.2), y + Inches(0.3), Inches(2.43), Inches(4.3))
        btf = box.text_frame
        btf.word_wrap = True

        p = btf.paragraphs[0]
        p.text = tag
        p.font.name = "Georgia"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(6)

        p2 = btf.add_paragraph()
        p2.text = head
        p2.font.name = "Calibri"
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_after = Pt(12)

        p3 = btf.add_paragraph()
        p3.text = desc
        p3.font.name = "Calibri"
        p3.font.size = Pt(12)
        p3.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 3: System Architecture (4 Functional Tiers)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Functional Architecture: The 4-Tier Operating Model", "SYSTEM ARCHITECTURE")

    tiers = [
        ("TIER 1: DELEGATE INTERFACE (iOS & Android)", 
         "VIP User Touchpoint", 
         "Dynamic Digital Pass (Apple/Google Wallet) • 5-Day Interactive Itinerary • 1-on-1 Meeting Scheduler • AI Face-Recognition Photo Gallery • Digital Business Card.",
         GOLD),
        ("TIER 2: CORE OPERATIONAL SERVICES", 
         "Business Logic & Workflows", 
         "Parliament Security Accreditation Engine • Inter-City Transit Scheduling • Multi-Currency Gateway & B2B Invoicing • Travel & Visa Document Vault.",
         WHITE),
        ("TIER 3: STAKEHOLDER & VENUE INTEGRATION", 
         "Institutional Interfaces", 
         "House of Commons Security Manifest Interface • University of Oxford & Cambridge Session Coordination • Gala Venue VIP Table Management.",
         GOLD_LIGHT),
        ("TIER 4: SECRETARIAT COMMAND CENTER", 
         "Organiser Control Center", 
         "Gate Check-In Mobile Scanner App (<2s access) • Real-Time Attendance Monitoring • Segmented VIP Push Broadcast • 24/7 VIP Concierge Hotline.",
         WHITE)
    ]

    for idx, (t_title, t_sub, t_desc, t_col) in enumerate(tiers):
        ly = Inches(1.8 + idx * 1.25)
        add_card(s3, Inches(0.8), ly, Inches(11.733), Inches(1.15), NAVY_CARD, GOLD if idx==0 else None)

        tbox = s3.shapes.add_textbox(Inches(1.0), ly + Inches(0.12), Inches(11.333), Inches(0.95))
        tf = tbox.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = t_title + "  |  " + t_sub
        p1.font.name = "Georgia"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = t_col
        p1.space_after = Pt(3)

        p2 = tf.add_paragraph()
        p2.text = t_desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 4: Delegate Credentialing & Security
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Delegate Credentialing & Parliamentary Security", "FEATURE SPECIFICATION")

    cards_s4 = [
        ("Accreditation Portal", "Pre-event identity capture formatted to satisfy British Parliamentary security vetting. Verifies legal name matching passport, nationality, and photo.", GOLD),
        ("Live Clearance Tracker", "Real-time status indicators informing delegates when their accreditation moves from 'Submitted' to 'Police Cleared' and 'Pass Issued'.", WHITE),
        ("Dynamic Anti-Fraud Pass", "Time-synchronized revolving QR code that prevents screenshot sharing and guarantees verified identity at venue gates.", GOLD_LIGHT),
        ("Apple & Google Wallet", "One-tap pass installation directly to iPhone or Android lock screens, configured to automatically appear near Westminster.", WHITE),
        ("Credential Tiers", "Visual badges distinguishing Award Winners, Executive Delegates, Keynote Speakers, Advisory Committee, and Media.", GOLD),
        ("Visa & Document Vault", "Secure offline repository for official stamped IBN UK Visa Support Letters, hotel confirmations, and transfer vouchers.", WHITE),
    ]

    for m, (f_title, f_desc, f_col) in enumerate(cards_s4):
        row = m // 3
        col = m % 3
        bx = Inches(0.8 + col * 4.04)
        by = Inches(1.8 + row * 2.5)
        add_card(s4, bx, by, Inches(3.64), Inches(2.2), NAVY_CARD)

        box = s4.shapes.add_textbox(bx + Inches(0.2), by + Inches(0.2), Inches(3.24), Inches(1.8))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "✦ " + f_title
        p1.font.name = "Georgia"
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = f_col
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 5: 5-Day Event & Multi-City Logistics
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "5-Day Event Logistics & Multi-City Navigation", "FEATURE SPECIFICATION")

    days_s5 = [
        ("DAY 1: 5 NOV", "House of Commons", "Awards & Reception", "Cromwell Green gate guidance, security protocol check-in, and formal Parliamentary Awards agenda.", GOLD),
        ("DAY 2: 6 NOV", "Central London", "Black Tie Gala", "VIP table allocations, multi-course dining timeline, and Black Tie / Tuxedo attire guidance.", WHITE),
        ("DAY 3: 7 NOV", "Univ of Oxford", "Leadership Forum", "Oxford coach departure tracking, keynote panel schedule, and live interactive session Q&A/polls.", GOLD_LIGHT),
        ("DAY 4: 8 NOV", "Cambridge", "Campus & Deep-Tech", "Cambridge departure alerts, deep-tech research hub itinerary, and collegiate walking route guides.", WHITE),
        ("DAY 5: 9 NOV", "Imperial College", "Innovation Showcase", "AI venture presentations, enterprise showcase, networking luncheon, and official closing delegation.", GOLD),
    ]

    for n, (d_num, d_city, d_event, d_desc, d_col) in enumerate(days_s5):
        dx = Inches(0.8 + n * 2.41)
        dy = Inches(1.8)
        add_card(s5, dx, dy, Inches(2.26), Inches(4.9), NAVY_CARD, GOLD if n==0 else None)

        box = s5.shapes.add_textbox(dx + Inches(0.18), dy + Inches(0.2), Inches(1.9), Inches(4.5))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = d_num
        p1.font.name = "Georgia"
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = d_col

        p2 = tf.add_paragraph()
        p2.text = d_city
        p2.font.name = "Calibri"
        p2.font.size = Pt(12)
        p2.font.color.rgb = MUTED_GRAY
        p2.space_after = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = d_event
        p3.font.name = "Calibri"
        p3.font.size = Pt(14)
        p3.font.bold = True
        p3.font.color.rgb = WHITE
        p3.space_after = Pt(12)

        p4 = tf.add_paragraph()
        p4.text = d_desc
        p4.font.name = "Calibri"
        p4.font.size = Pt(11)
        p4.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 6: Financial Management & B2B Invoicing
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Financial Management & B2B Invoicing Hub", "FEATURE SPECIFICATION")

    cards_s6 = [
        ("Multi-Currency Checkout", "Seamless checkout supporting British Pounds (£), US Dollars ($), Euros (€), and Indian Rupees (₹). Supports international corporate credit cards, Apple Pay, Google Pay, UPI, and Net Banking.", GOLD),
        ("Automated B2B Tax Invoicing", "Instant PDF generation of compliant commercial tax invoices including Corporate Name, GSTIN, and UK VAT details for corporate expense claims.", WHITE),
        ("Regulatory Package Transparency", "Itemizes packages exclusively for the 'IBN Business Leadership Programme' (academic sessions, gala hospitality, and transit), maintaining parliamentary non-commercial status.", GOLD_LIGHT),
        ("Wire Transfer / SWIFT Flow", "Enables corporate finance teams to upload official bank transfer or SWIFT/NEFT remittance receipts with automated secretariat verification.", WHITE),
    ]

    for q, (f_head, f_body, f_col) in enumerate(cards_s6):
        row = q // 2
        col = q % 2
        fx = Inches(0.8 + col * 6.04)
        fy = Inches(1.8 + row * 2.5)
        add_card(s6, fx, fy, Inches(5.68), Inches(2.2), NAVY_CARD, GOLD if q==0 else None)

        box = s6.shapes.add_textbox(fx + Inches(0.3), fy + Inches(0.25), Inches(5.08), Inches(1.7))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "✦ " + f_head
        p1.font.name = "Georgia"
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = f_col
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = f_body
        p2.font.name = "Calibri"
        p2.font.size = Pt(12)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 7: Networking & Meeting Matchmaker
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "Executive Networking & Business Matchmaking", "FEATURE SPECIFICATION")

    cards_s7 = [
        ("Verified 100-Leader Directory", "Private, searchable roster of attending CEOs, investors, and founders. Filterable by industry sector, geographical focus, and investment appetite.", GOLD),
        ("1-on-1 Meeting Scheduler", "Empowers delegates to request, schedule, and confirm 15-minute private meetings during designated summit networking breaks with built-in agenda sync.", WHITE),
        ("Privacy-Protected In-App Chat", "Facilitates direct executive messaging within the platform. Mobile numbers and personal email addresses remain shielded unless mutually shared.", GOLD_LIGHT),
        ("Digital Business Card Exchange", "Instant peer contact exchange via contactless phone tap (NFC) or badge QR scan, saving directly into mobile phone address books.", WHITE),
    ]

    for r, (n_head, n_body, n_col) in enumerate(cards_s7):
        row = r // 2
        col = r % 2
        nx = Inches(0.8 + col * 6.04)
        ny = Inches(1.8 + row * 2.5)
        add_card(s7, nx, ny, Inches(5.68), Inches(2.2), NAVY_CARD, GOLD if r==1 else None)

        box = s7.shapes.add_textbox(nx + Inches(0.3), ny + Inches(0.25), Inches(5.08), Inches(1.7))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "✦ " + n_head
        p1.font.name = "Georgia"
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = n_col
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = n_body
        p2.font.name = "Calibri"
        p2.font.size = Pt(12)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 8: AI Photo Experience & PR
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "AI-Powered Photo Experience & Media Hub", "FEATURE SPECIFICATION")

    l_card = add_card(s8, Inches(0.8), Inches(1.8), Inches(7.2), Inches(4.9), NAVY_CARD, GOLD)
    lbox = s8.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(6.6), Inches(4.3))
    ltf = lbox.text_frame
    ltf.word_wrap = True

    lp1 = ltf.paragraphs[0]
    lp1.text = "SEAMLESS PERSONALIZED MEDIA DISTRIBUTION"
    lp1.font.name = "Georgia"
    lp1.font.size = Pt(15)
    lp1.font.bold = True
    lp1.font.color.rgb = GOLD
    lp1.space_after = Pt(12)

    steps_ai = [
        ("1. Photographer Ingestion", "Official summit media teams upload thousands of high-res event photos directly from the House of Commons, Gala, and Universities."),
        ("2. AI Facial Recognition", "Delegates upload a single reference selfie during onboarding. The platform automatically indexes and isolates all images containing that delegate."),
        ("3. Private 'My Moments' Gallery", "Delegates access a personalized, curated gallery containing exclusively their summit moments within hours of session completion."),
        ("4. High-Res PR Export", "Single-tap download of press-ready, high-resolution RAW/JPEG assets ready for corporate communications and immediate LinkedIn sharing.")
    ]
    for st_title, st_desc in steps_ai:
        p_st = ltf.add_paragraph()
        p_st.text = f"• {st_title}: {st_desc}"
        p_st.font.name = "Calibri"
        p_st.font.size = Pt(12)
        p_st.font.color.rgb = WHITE
        p_st.space_after = Pt(8)

    r_stats = [
        ("INSTANT ACCESS", "Zero waiting for manual photographer sorting"),
        ("PRESS-READY", "Full-resolution original print quality"),
        ("1-CLICK SOCIAL", "Pre-formatted exports for LinkedIn & Instagram"),
    ]
    for s_idx, (s_title, s_sub) in enumerate(r_stats):
        sy = Inches(1.8 + s_idx * 1.68)
        add_card(s8, Inches(8.3), sy, Inches(4.233), Inches(1.5), NAVY_LIGHT)
        sbox = s8.shapes.add_textbox(Inches(8.5), sy + Inches(0.2), Inches(3.8), Inches(1.1))
        stf = sbox.text_frame
        stf.word_wrap = True

        p_st1 = stf.paragraphs[0]
        p_st1.text = s_title
        p_st1.font.name = "Georgia"
        p_st1.font.size = Pt(15)
        p_st1.font.bold = True
        p_st1.font.color.rgb = GOLD

        p_st2 = stf.add_paragraph()
        p_st2.text = s_sub
        p_st2.font.name = "Calibri"
        p_st2.font.size = Pt(11.5)
        p_st2.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 9: Secretariat Command & Operations
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9)
    add_header(s9, "Onsite Secretariat Operations & Control Center", "OPERATIONAL ARCHITECTURE")

    cards_s9 = [
        ("Gate Check-in Scanner App", "Dedicated staff mode for mobile phones and tablets to scan delegate QR codes at the House of Commons and Gala doors in under two seconds.", GOLD),
        ("Live Headcount Dashboard", "Real-time visibility into attendance counts inside the House of Commons, banquet halls, and passenger boarding for inter-city coach transfers.", WHITE),
        ("Parliament Manifest Export", "One-click export of the accredited guest manifest conforming exactly to British Parliamentary Security and Police format requirements.", GOLD_LIGHT),
        ("Emergency VIP Concierge", "Integrated one-touch hotline connecting delegates directly to the IBN Event Coordinator and secretariat leads for instant travel and venue support.", WHITE),
    ]

    for v, (o_head, o_body, o_col) in enumerate(cards_s9):
        row = v // 2
        col = v % 2
        ox = Inches(0.8 + col * 6.04)
        oy = Inches(1.8 + row * 2.5)
        add_card(s9, ox, oy, Inches(5.68), Inches(2.2), NAVY_CARD, GOLD if v==0 else None)

        box = s9.shapes.add_textbox(ox + Inches(0.3), oy + Inches(0.25), Inches(5.08), Inches(1.7))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "✦ " + o_head
        p1.font.name = "Georgia"
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = o_col
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = o_body
        p2.font.name = "Calibri"
        p2.font.size = Pt(12)
        p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 10: Lifecycle & Rollout Roadmap
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10)
    add_header(s10, "End-to-End Delegate Lifecycle & Delivery Roadmap", "DEPLOYMENT ARCHITECTURE")

    stages_s10 = [
        ("STAGE 1: PRE-EVENT", "Registration & KYC", "Profile registration, passport data submission, UK visa letter issuance, and package payment processing with automated B2B invoicing."),
        ("STAGE 2: ACCREDITATION", "Parliament Clearance", "Security vetting by authorities, approval notification, and activation of Apple/Google Wallet dynamic passes."),
        ("STAGE 3: LIVE EVENT", "5-Day Summit Execution", "Fast-track QR entry, coach transit navigation, session Q&A participation, and 1-on-1 meeting execution across London, Oxford & Cambridge."),
        ("STAGE 4: POST-EVENT", "Media & Continuity", "Instant AI photo vault downloads, certificate distribution, and ongoing bilateral directory networking."),
    ]

    for t_idx, (st_phase, st_name, st_det) in enumerate(stages_s10):
        tx = Inches(0.8 + t_idx * 3.03)
        ty = Inches(1.8)
        add_card(s10, tx, ty, Inches(2.83), Inches(4.9), NAVY_CARD, GOLD if t_idx==2 else None)

        box = s10.shapes.add_textbox(tx + Inches(0.2), ty + Inches(0.3), Inches(2.43), Inches(4.3))
        tf = box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = st_phase
        p1.font.name = "Georgia"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = GOLD
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = st_name
        p2.font.name = "Calibri"
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_after = Pt(12)

        p3 = tf.add_paragraph()
        p3.text = st_det
        p3.font.name = "Calibri"
        p3.font.size = Pt(12)
        p3.font.color.rgb = MUTED_GRAY

    output_path = "IBN_2026_Event_and_Mobile_App_Executive_Presentation.pptx"
    prs.save(output_path)
    print(f"Master presentation saved successfully to {output_path}")

if __name__ == "__main__":
    build_master_deck()
