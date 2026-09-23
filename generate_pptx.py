import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # Set 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    # Colors
    NAVY_DARK = RGBColor(6, 18, 38)       # #061226
    NAVY_CARD = RGBColor(12, 28, 54)      # #0C1C36
    NAVY_LIGHT = RGBColor(20, 42, 78)     # #142A4E
    GOLD = RGBColor(212, 175, 55)         # #D4AF37
    GOLD_LIGHT = RGBColor(243, 229, 171)  # #F3E5AB
    WHITE = RGBColor(255, 255, 255)
    MUTED_GRAY = RGBColor(180, 195, 215)
    GREEN_ACCENT = RGBColor(16, 124, 65)

    def set_slide_background(slide, color=NAVY_DARK):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="GLOBAL INDIAN BUSINESS EXCELLENCE AWARDS 2026"):
        # Header banner text
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Georgia"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = GOLD
        p_cat.space_after = Pt(4)

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.name = "Calibri"
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        # Gold accent line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(1.8), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = GOLD
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=NAVY_CARD, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.color.rgb = RGBColor(30, 52, 88)
            card.line.width = Pt(1)
        return card

    # ==========================================
    # SLIDE 1: Title & Cover Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, NAVY_DARK)

    # Decorative border
    outer_box = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.4), Inches(12.533), Inches(6.7))
    outer_box.fill.background()
    outer_box.line.color.rgb = RGBColor(40, 70, 110)
    outer_box.line.width = Pt(1)

    inner_box = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.45), Inches(12.433), Inches(6.6))
    inner_box.fill.background()
    inner_box.line.color.rgb = GOLD
    inner_box.line.width = Pt(0.75)

    # Top presentation credit
    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(0.8))
    tf = tbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "INDIAN BUSINESS NETWORK"
    p.font.name = "Georgia"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GOLD

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Connect • Collaborate • Inspire"
    p2.font.name = "Calibri"
    p2.font.size = Pt(11)
    p2.font.color.rgb = MUTED_GRAY

    # Main Title
    title_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(2.2))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    p_main = tf2.paragraphs[0]
    p_main.alignment = PP_ALIGN.CENTER
    p_main.text = "GLOBAL INDIAN\nBUSINESS EXCELLENCE AWARDS 2026"
    p_main.font.name = "Georgia"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = GOLD_LIGHT
    p_main.space_after = Pt(12)

    p_sub = tf2.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.text = "Honouring Indian Business Leaders Creating Global Impact"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = WHITE

    # Venue & Dates Card
    venue_card = add_card(s1, Inches(2.5), Inches(4.2), Inches(8.333), Inches(1.5), NAVY_CARD, GOLD)
    vbox = s1.shapes.add_textbox(Inches(2.7), Inches(4.35), Inches(7.933), Inches(1.2))
    vtf = vbox.text_frame
    vp1 = vtf.paragraphs[0]
    vp1.alignment = PP_ALIGN.CENTER
    vp1.text = "🏛️ HOUSE OF COMMONS, BRITISH PARLIAMENT | LONDON, UK"
    vp1.font.name = "Calibri"
    vp1.font.size = Pt(16)
    vp1.font.bold = True
    vp1.font.color.rgb = GOLD

    vp2 = vtf.add_paragraph()
    vp2.alignment = PP_ALIGN.CENTER
    vp2.text = "📅 5–9 NOVEMBER 2026  •  🔒 BY INVITATION ONLY (STRICTLY CAPPED AT 100 GUESTS)"
    vp2.font.name = "Calibri"
    vp2.font.size = Pt(13)
    vp2.font.bold = True
    vp2.font.color.rgb = WHITE
    vp2.space_before = Pt(6)

    # Footer note
    ft_box = s1.shapes.add_textbox(Inches(1.0), Inches(6.1), Inches(11.333), Inches(0.6))
    ft_tf = ft_box.text_frame
    fp = ft_tf.paragraphs[0]
    fp.alignment = PP_ALIGN.CENTER
    fp.text = "CEOs  •  ENTREPRENEURS  •  INVESTORS  •  GLOBAL PROFESSIONALS"
    fp.font.name = "Calibri"
    fp.font.size = Pt(12)
    fp.font.bold = True
    fp.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 2: Executive Vision & Core Pillars
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, NAVY_DARK)
    add_header(s2, "Executive Vision: Redefining Global Indian Impact")

    pillars = [
        ("01. HONOUR", "Recognizing Excellence", "Celebrating transformative entrepreneurs, CEOs, and investors of Indian heritage who have built sustainable cross-border enterprises and global market leadership.", GOLD),
        ("02. CONNECT", "Bespoke Bilateral Network", "An intimate, invitation-only congregation of 100 distinguished leaders, facilitating direct engagement between Indian titans and UK political and business dignitaries.", WHITE),
        ("03. INSPIRE", "Academic Immersion", "Forging institutional bridges with Oxford, Cambridge, and Imperial College London to exchange cutting-edge research, venture creation, and governance best practices.", GOLD_LIGHT),
    ]

    card_w = Inches(3.64)
    card_h = Inches(4.8)
    for i, (tag, heading, body, accent_color) in enumerate(pillars):
        x = Inches(0.8 + i * 4.04)
        y = Inches(1.8)
        add_card(s2, x, y, card_w, card_h, NAVY_CARD, RGBColor(40, 70, 110) if i!=0 else GOLD)

        t_box = s2.shapes.add_textbox(x + Inches(0.3), y + Inches(0.4), card_w - Inches(0.6), card_h - Inches(0.8))
        tf = t_box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = tag
        p1.font.name = "Georgia"
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = accent_color
        p1.space_after = Pt(10)

        p2 = tf.add_paragraph()
        p2.text = heading
        p2.font.name = "Calibri"
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_after = Pt(14)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.name = "Calibri"
        p3.font.size = Pt(13)
        p3.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 3: The Historic Venue: House of Commons
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, NAVY_DARK)
    add_header(s3, "The Historic Stage: House of Commons, British Parliament")

    # Left Column: Detailed Venue narrative
    left_card = add_card(s3, Inches(0.8), Inches(1.8), Inches(7.2), Inches(4.9), NAVY_CARD)
    lt_box = s3.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(6.4), Inches(4.3))
    ltf = lt_box.text_frame
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = "AN UNPARALLELED ASSEMBLY AT WESTMINSTER"
    p.font.name = "Georgia"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p.space_after = Pt(10)

    bullets = [
        "Palace of Westminster Heritage: Gathering at the epicenter of British history and global democratic tradition.",
        "High-Level India–UK Bilateral Synergy: Positioned as an executive bridge between leading Indian business houses and UK political, trade, and economic leaders.",
        "Unrivalled Exclusivity: Strictly capped at 100 vetted delegates to maintain intimate networking and highest security protocols.",
        "Parliamentary Reception: Formal awards presentation followed by an executive networking reception inside the parliamentary estate.",
    ]
    for b in bullets:
        bp = ltf.add_paragraph()
        bp.text = "• " + b
        bp.font.name = "Calibri"
        bp.font.size = Pt(13)
        bp.font.color.rgb = WHITE
        bp.space_after = Pt(8)

    # Right Column: Key Metrics / Stats
    stats = [
        ("100", "DISTINGUISHED GUESTS", "Exclusive C-Suite audience cap"),
        ("5 DAYS", "EXPANDED PROGRAMME", "London • Oxford • Cambridge"),
        ("3 HUBS", "WORLD TOP UNIVERSITIES", "Oxford, Cambridge & Imperial"),
    ]
    for j, (num, lbl, desc) in enumerate(stats):
        sy = Inches(1.8 + j * 1.68)
        scard = add_card(s3, Inches(8.3), sy, Inches(4.2), Inches(1.5), NAVY_LIGHT, GOLD if j==0 else None)
        sbox = s3.shapes.add_textbox(Inches(8.5), sy + Inches(0.15), Inches(3.8), Inches(1.2))
        stf = sbox.text_frame
        stf.word_wrap = True
        
        sp1 = stf.paragraphs[0]
        sp1.text = num
        sp1.font.name = "Georgia"
        sp1.font.size = Pt(28)
        sp1.font.bold = True
        sp1.font.color.rgb = GOLD
        
        sp2 = stf.add_paragraph()
        sp2.text = lbl + " — " + desc
        sp2.font.name = "Calibri"
        sp2.font.size = Pt(11)
        sp2.font.bold = True
        sp2.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 4: 5-Day Master Programme Timeline
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, NAVY_DARK)
    add_header(s4, "5-Day Executive Programme Itinerary (5–9 Nov 2026)")

    days = [
        ("DAY 1", "5 NOV 2026", "House of Commons", "Global Indian Business Excellence Awards & Executive Networking Reception at the British Parliament.", GOLD),
        ("DAY 2", "6 NOV 2026", "Central London", "Black Tie Gala Dinner & Executive Networking: Gourmet dining, bilateral trade addresses & strategic dialogue.", WHITE),
        ("DAY 3", "7 NOV 2026", "Univ of Oxford", "Oxford Business Leadership Forum & Panel Discussions on global enterprise, tech, and economic policy.", GOLD_LIGHT),
        ("DAY 4", "8 NOV 2026", "Univ of Cambridge", "University of Cambridge Campus Visit: Deep-tech incubation, spinouts & historic academic immersion.", WHITE),
        ("DAY 5", "9 NOV 2026", "Imperial College", "Imperial College London Innovation & Research Experience: AI, frontier tech & entrepreneurship showcase.", GOLD),
    ]

    day_w = Inches(2.26)
    day_h = Inches(4.9)
    for k, (d_num, d_date, d_loc, d_text, d_accent) in enumerate(days):
        dx = Inches(0.8 + k * 2.41)
        dy = Inches(1.8)
        add_card(s4, dx, dy, day_w, day_h, NAVY_CARD, RGBColor(40, 70, 110) if k!=0 else GOLD)

        dt_box = s4.shapes.add_textbox(dx + Inches(0.18), dy + Inches(0.25), day_w - Inches(0.36), day_h - Inches(0.5))
        dtf = dt_box.text_frame
        dtf.word_wrap = True

        dp1 = dtf.paragraphs[0]
        dp1.text = d_num
        dp1.font.name = "Georgia"
        dp1.font.size = Pt(16)
        dp1.font.bold = True
        dp1.font.color.rgb = d_accent
        dp1.space_after = Pt(2)

        dp2 = dtf.add_paragraph()
        dp2.text = d_date
        dp2.font.name = "Calibri"
        dp2.font.size = Pt(12)
        dp2.font.bold = True
        dp2.font.color.rgb = MUTED_GRAY
        dp2.space_after = Pt(8)

        dp3 = dtf.add_paragraph()
        dp3.text = "📍 " + d_loc
        dp3.font.name = "Calibri"
        dp3.font.size = Pt(13)
        dp3.font.bold = True
        dp3.font.color.rgb = WHITE
        dp3.space_after = Pt(10)

        dp4 = dtf.add_paragraph()
        dp4.text = d_text
        dp4.font.name = "Calibri"
        dp4.font.size = Pt(11)
        dp4.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 5: Academic & Research Immersion
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, NAVY_DARK)
    add_header(s5, "Beyond London: Oxford, Cambridge & Imperial College")

    acad = [
        ("UNIVERSITY OF OXFORD", "Oxford Business Leadership Forum", "7 November 2026", 
         ["High-level panel debates on global trade & geopolitics", 
          "Executive interaction with Oxford scholars & fellows", 
          "Historic collegiate walking tour & private dining options"]),
        ("UNIVERSITY OF CAMBRIDGE", "Deep-Tech & Campus Immersion", "8 November 2026", 
         ["Cambridge Cluster & enterprise spinout ecosystem", 
          "Bridging academic scientific research with capital", 
          "Exclusive tour of historic Cambridge collegiate grounds"]),
        ("IMPERIAL COLLEGE LONDON", "Innovation & Research Showcase", "9 November 2026", 
         ["Focus on Artificial Intelligence, MedTech & CleanTech", 
          "Executive briefing at Imperial Enterprise Hub", 
          "Commercialisation avenues for Indo-UK partnerships"])
    ]

    for m, (u_title, u_sub, u_date, u_bullets) in enumerate(acad):
        ux = Inches(0.8 + m * 4.04)
        uy = Inches(1.8)
        add_card(s5, ux, uy, Inches(3.64), Inches(4.9), NAVY_CARD, GOLD if m==0 else None)

        ubox = s5.shapes.add_textbox(ux + Inches(0.25), uy + Inches(0.3), Inches(3.14), Inches(4.3))
        utf = ubox.text_frame
        utf.word_wrap = True

        up1 = utf.paragraphs[0]
        up1.text = u_title
        up1.font.name = "Georgia"
        up1.font.size = Pt(13)
        up1.font.bold = True
        up1.font.color.rgb = GOLD
        up1.space_after = Pt(4)

        up2 = utf.add_paragraph()
        up2.text = u_sub
        up2.font.name = "Calibri"
        up2.font.size = Pt(15)
        up2.font.bold = True
        up2.font.color.rgb = WHITE

        up_date = utf.add_paragraph()
        up_date.text = "📅 " + u_date
        up_date.font.name = "Calibri"
        up_date.font.size = Pt(11)
        up_date.font.color.rgb = GOLD_LIGHT
        up_date.space_after = Pt(12)

        for ub in u_bullets:
            p_b = utf.add_paragraph()
            p_b.text = "• " + ub
            p_b.font.name = "Calibri"
            p_b.font.size = Pt(12)
            p_b.font.color.rgb = MUTED_GRAY
            p_b.space_after = Pt(6)

    # ==========================================
    # SLIDE 6: Executive Delegate Inclusions
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, NAVY_DARK)
    add_header(s6, "Executive Delegate Package & Deliverables")

    inclusions_list = [
        ("Executive Delegate Badge", "Accredited credentials granting full entry to official sessions, forums, and private transports.", GOLD),
        ("Parliamentary Awards Access", "Official organiser nomination for entry into the House of Commons Awards & Reception.", WHITE),
        ("Black Tie Gala Dinner", "Seat at the prestigious central London Gala Dinner with gourmet multi-course hospitality.", GOLD_LIGHT),
        ("Academic Forums Pass", "Full delegation access to University of Oxford, Cambridge, and Imperial College experiences.", WHITE),
        ("Delegate Welcome Kit", "Curated luxury delegate stationery, commemorative executive mementos & summit briefcase.", GOLD),
        ("Professional Event Photography", "Complete personal & corporate high-resolution media package by official UK press photographers.", WHITE),
    ]

    for n, (inc_title, inc_desc, inc_col) in enumerate(inclusions_list):
        row = n // 3
        col = n % 3
        ix = Inches(0.8 + col * 4.04)
        iy = Inches(1.8 + row * 2.5)
        add_card(s6, ix, iy, Inches(3.64), Inches(2.2), NAVY_CARD)

        ibox = s6.shapes.add_textbox(ix + Inches(0.25), iy + Inches(0.25), Inches(3.14), Inches(1.7))
        itf = ibox.text_frame
        itf.word_wrap = True

        ip1 = itf.paragraphs[0]
        ip1.text = "✦ " + inc_title
        ip1.font.name = "Georgia"
        ip1.font.size = Pt(15)
        ip1.font.bold = True
        ip1.font.color.rgb = inc_col
        ip1.space_after = Pt(8)

        ip2 = itf.add_paragraph()
        ip2.text = inc_desc
        ip2.font.name = "Calibri"
        ip2.font.size = Pt(12)
        ip2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 7: Award Categories & Recognition
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, NAVY_DARK)
    add_header(s7, "Award Categories & Selection Governance")

    cats = [
        "Global Business Leader of the Year",
        "Disruptive Tech Innovator Award",
        "Excellence in Cross-Border Trade & Manufacturing",
        "Sustainable Enterprise & ESG Leadership",
        "Healthcare & Life Sciences Pioneer",
        "Next-Gen Entrepreneur of the Year",
        "Financial Services & Investment Excellence",
        "Women in Global Business Leadership"
    ]

    # Left column: Categories List
    cat_card = add_card(s7, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.9), NAVY_CARD)
    cbox = s7.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(6.2), Inches(4.3))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    cp1 = ctf.paragraphs[0]
    cp1.text = "CORE RECOGNITION CATEGORIES"
    cp1.font.name = "Georgia"
    cp1.font.size = Pt(14)
    cp1.font.bold = True
    cp1.font.color.rgb = GOLD
    cp1.space_after = Pt(10)

    for c in cats:
        cp = ctf.add_paragraph()
        cp.text = "🏆  " + c
        cp.font.name = "Calibri"
        cp.font.size = Pt(12.5)
        cp.font.bold = True
        cp.font.color.rgb = WHITE
        cp.space_after = Pt(5)

    # Right column: Selection criteria
    gov_card = add_card(s7, Inches(7.9), Inches(1.8), Inches(4.6), Inches(4.9), NAVY_CARD, GOLD)
    gbox = s7.shapes.add_textbox(Inches(8.2), Inches(2.1), Inches(4.0), Inches(4.3))
    gtf = gbox.text_frame
    gtf.word_wrap = True

    gp1 = gtf.paragraphs[0]
    gp1.text = "EVALUATION CRITERIA"
    gp1.font.name = "Georgia"
    gp1.font.size = Pt(14)
    gp1.font.bold = True
    gp1.font.color.rgb = GOLD
    gp1.space_after = Pt(10)

    eval_points = [
        "Global Footprint & Scalability: Demonstrated cross-border operations and international client reach.",
        "Innovation & Disruption: Pioneering IP, technology adoption, or distinctive business models.",
        "Integrity & Corporate Governance: Exemplary ethical standing, leadership resilience, and ESG compliance.",
        "Peer & Advisory Vetting: Evaluated by the IBN Advisory Committee and external industry luminaries."
    ]
    for ep in eval_points:
        e_par = gtf.add_paragraph()
        e_par.text = "✔ " + ep
        e_par.font.name = "Calibri"
        e_par.font.size = Pt(12)
        e_par.font.color.rgb = MUTED_GRAY
        e_par.space_after = Pt(8)

    # ==========================================
    # SLIDE 8: Governance, Protocols & Compliance
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, NAVY_DARK)
    add_header(s8, "Institutional Compliance & Security Protocols")

    # Important note banner
    banner = add_card(s8, Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.1), NAVY_CARD, GOLD)
    bbox = s8.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(1.7))
    btf = bbox.text_frame
    btf.word_wrap = True

    bp1 = btf.paragraphs[0]
    bp1.text = "OFFICIAL COMPLIANCE & LEGAL NOTICE"
    bp1.font.name = "Georgia"
    bp1.font.size = Pt(14)
    bp1.font.bold = True
    bp1.font.color.rgb = GOLD
    bp1.space_after = Pt(6)

    bp2 = btf.add_paragraph()
    bp2.text = (
        "The Global Indian Business Excellence Awards 2026 at the House of Commons, British Parliament, London, "
        "is an invitation-only event. Invitations are extended at the sole discretion of the organisers and are subject "
        "to venue protocols and capacity. No fee is charged for attendance at the Awards Ceremony, and the invitation is "
        "not sold or assigned any monetary value. Delegate fees relate exclusively to the curated IBN Business Leadership Programme, "
        "academic forums at Oxford, Cambridge & Imperial, and associated networking activities."
    )
    bp2.font.name = "Calibri"
    bp2.font.size = Pt(12)
    bp2.font.color.rgb = WHITE

    # Two detail cards below
    c1 = add_card(s8, Inches(0.8), Inches(4.2), Inches(5.7), Inches(2.5), NAVY_CARD)
    c1_box = s8.shapes.add_textbox(Inches(1.0), Inches(4.4), Inches(5.3), Inches(2.1))
    c1_tf = c1_box.text_frame
    c1_tf.word_wrap = True
    c1_p1 = c1_tf.paragraphs[0]
    c1_p1.text = "PARLIAMENT SECURITY PROTOCOL"
    c1_p1.font.name = "Calibri"
    c1_p1.font.size = Pt(14)
    c1_p1.font.bold = True
    c1_p1.font.color.rgb = GOLD_LIGHT
    c1_p1.space_after = Pt(6)
    c1_p2 = c1_tf.add_paragraph()
    c1_p2.text = "• Formal photo ID (valid Passport) mandatory at Cromwell Green entrance.\n• Advance security clearance required 30 days prior to the summit.\n• Dress Code: Business Formal / Parliamentary Lounge Suit / Formal Indian."
    c1_p2.font.name = "Calibri"
    c1_p2.font.size = Pt(11.5)
    c1_p2.font.color.rgb = MUTED_GRAY

    c2 = add_card(s8, Inches(6.8), Inches(4.2), Inches(5.7), Inches(2.5), NAVY_CARD)
    c2_box = s8.shapes.add_textbox(Inches(7.0), Inches(4.4), Inches(5.3), Inches(2.1))
    c2_tf = c2_box.text_frame
    c2_tf.word_wrap = True
    c2_p1 = c2_tf.paragraphs[0]
    c2_p1.text = "UK VISA & LOGISTICS CONCIERGE"
    c2_p1.font.name = "Calibri"
    c2_p1.font.size = Pt(14)
    c2_p1.font.bold = True
    c2_p1.font.color.rgb = GOLD_LIGHT
    c2_p1.space_after = Pt(6)
    c2_p2 = c2_tf.add_paragraph()
    c2_p2.text = "• Official IBN Visa Support Letters issued upon registration confirmation.\n• Dedicated coach transfers between Central London, Oxford & Cambridge.\n• Partner 5-star hotel preferential booking links provided to confirmed delegates."
    c2_p2.font.name = "Calibri"
    c2_p2.font.size = Pt(11.5)
    c2_p2.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 9: Leadership Committee & Secretariat
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, NAVY_DARK)
    add_header(s9, "Organising Leadership & Secretariat Contacts")

    leaders_info = [
        ("VELOU SINGARAM", "President", "Indian Business Network", "Overseeing bilateral trade delegations and presidential addresses."),
        ("JAYABALAN", "Secretary", "Tel: +44 7960 446339", "Secretariat operations, parliamentary liaison & institutional protocol."),
        ("NAGARAJAN", "Treasurer", "Tel: +44 7383 969604", "Financial governance, corporate sponsorships & delegate accounts."),
        ("SUBHA AUSTALEKSHMI", "Event Coordinator", "Tel: +44 7587 260254", "Delegate onboarding, accommodation, visa letters & concierge."),
    ]

    l_w = Inches(2.78)
    l_h = Inches(4.8)
    for p_idx, (l_name, l_role, l_contact, l_desc) in enumerate(leaders_info):
        lx = Inches(0.8 + p_idx * 2.97)
        ly = Inches(1.8)
        add_card(s9, lx, ly, l_w, l_h, NAVY_CARD, GOLD if p_idx==0 else None)

        lbox = s9.shapes.add_textbox(lx + Inches(0.2), ly + Inches(0.4), l_w - Inches(0.4), l_h - Inches(0.8))
        ltf = lbox.text_frame
        ltf.word_wrap = True

        lp1 = ltf.paragraphs[0]
        lp1.text = l_name
        lp1.font.name = "Georgia"
        lp1.font.size = Pt(13)
        lp1.font.bold = True
        lp1.font.color.rgb = GOLD
        lp1.space_after = Pt(2)

        lp2 = ltf.add_paragraph()
        lp2.text = l_role
        lp2.font.name = "Calibri"
        lp2.font.size = Pt(12)
        lp2.font.bold = True
        lp2.font.color.rgb = WHITE
        lp2.space_after = Pt(4)

        lp3 = ltf.add_paragraph()
        lp3.text = l_contact
        lp3.font.name = "Calibri"
        lp3.font.size = Pt(11)
        lp3.font.color.rgb = GOLD_LIGHT
        lp3.space_after = Pt(12)

        lp4 = ltf.add_paragraph()
        lp4.text = l_desc
        lp4.font.name = "Calibri"
        lp4.font.size = Pt(11)
        lp4.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 10: Registration & Call to Action
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, NAVY_DARK)
    add_header(s10, "Confirm Your Delegate Participation")

    # Center card
    center_card = add_card(s10, Inches(1.8), Inches(1.8), Inches(9.733), Inches(4.8), NAVY_CARD, GOLD)
    cbox = s10.shapes.add_textbox(Inches(2.2), Inches(2.1), Inches(8.933), Inches(4.2))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    p_c1 = ctf.paragraphs[0]
    p_c1.alignment = PP_ALIGN.CENTER
    p_c1.text = "STRICTLY LIMITED TO 100 DISTINGUISHED GUESTS"
    p_c1.font.name = "Georgia"
    p_c1.font.size = Pt(18)
    p_c1.font.bold = True
    p_c1.font.color.rgb = GOLD
    p_c1.space_after = Pt(10)

    p_c2 = ctf.add_paragraph()
    p_c2.alignment = PP_ALIGN.CENTER
    p_c2.text = "Due to parliamentary accreditation requirements, registrations will close immediately upon reaching capacity."
    p_c2.font.name = "Calibri"
    p_c2.font.size = Pt(13)
    p_c2.font.color.rgb = MUTED_GRAY
    p_c2.space_after = Pt(18)

    # 3 Steps
    steps = [
        ("Step 1", "Submit Registration", "Access tally.so/r/NpzrWj to register your executive profile and company credentials."),
        ("Step 2", "Secretariat Review", "IBN Advisory Board reviews eligibility and initiates parliamentary clearance accreditation."),
        ("Step 3", "Official Delegate Pack", "Receive confirmation, visa invitation letter, and personal summit onboarding kit.")
    ]
    for s_step, s_title, s_desc in steps:
        p_step = ctf.add_paragraph()
        p_step.text = f"✦ {s_step}: {s_title} — {s_desc}"
        p_step.font.name = "Calibri"
        p_step.font.size = Pt(12)
        p_step.font.color.rgb = WHITE
        p_step.space_after = Pt(8)

    p_cta = ctf.add_paragraph()
    p_cta.alignment = PP_ALIGN.CENTER
    p_cta.text = "REGISTER TODAY: https://tally.so/r/NpzrWj\nDirect Inquiries: events@ibnonline.co.uk | +44 7587 260254"
    p_cta.font.name = "Calibri"
    p_cta.font.size = Pt(14)
    p_cta.font.bold = True
    p_cta.font.color.rgb = GOLD
    p_cta.space_before = Pt(14)

    output_path = "Global_Indian_Business_Excellence_Awards_2026.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
